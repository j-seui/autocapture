"""Opt-in native smoke test against a locally generated mock window."""
from __future__ import annotations

import json
import tkinter as tk
from pathlib import Path
from tkinter import ttk

from PIL import Image, ImageGrab
from pptx import Presentation

from .models import Config, Rect, Tab
from .progress import CaptureProgress
from .runner import Result, save_capture
from .win32 import Desktop, U, window_identity


def self_test(output_dir: str) -> None:
    output = Path(output_dir).resolve()
    output.mkdir(parents=True, exist_ok=True)
    previous_decks = set(output.glob("*.pptx"))
    control = tk.Tk()
    control.title("Self-test controller")
    control.geometry("430x140+50+50")
    ttk.Label(control, text="Bloomberg 캡처 모의 화면 검증").pack(pady=12)
    target = tk.Toplevel(control)
    target.title("Launchpad - MOCK TEST")
    target.geometry("800x70+100+100")
    target.configure(bg="#333333")
    header = tk.Frame(target, bg="#333333", height=40)
    header.pack(fill="x")
    graph = tk.Toplevel(control)
    graph.title("Independent graph window - MOCK TEST")
    region = tk.Frame(graph, bg="#ffffff")
    region.pack(fill="both", expand=True)

    def select(index: int) -> None:
        graph.geometry("780x500+100+250" if index == 0 else "550x360+100+250")
        region.configure(bg="#f4f6fa" if index == 0 else "#d6e9da")
        label.configure(text=("My Tab" if index == 0 else "자유 입력") + "\n\nMOCK CAPTURE\n한글·축·범례 영역 테스트")
        graph.lift()
        graph.focus_force()

    for index, name in enumerate(("My Tab", "자유 입력")):
        tk.Button(header, text=name, command=lambda i=index: select(i), width=12).pack(side="left", padx=3, pady=4)
    label = tk.Label(region, bg="white", fg="#222222", font=("Malgun Gothic", 22))
    label.pack(expand=True)
    select(0)
    control.update()
    hwnd = U.GetAncestor(target.winfo_id(), 2)
    desktop = Desktop(window_identity(hwnd))
    bounds, dpi = desktop.geometry()
    graph_desktop = Desktop(window_identity(U.GetAncestor(graph.winfo_id(), 2)))
    graph_bounds, _ = graph_desktop.geometry()
    layout = desktop.monitor_rects()
    settings = Config(target.title(), desktop.window.class_name, bounds.width, bounds.height, dpi, str(output), [
        Tab("My Tab", (45, 18), Rect(graph_bounds.x, graph_bounds.y, 780, 500), .7, "screen", layout),
        Tab("자유 입력", (155, 18), Rect(graph_bounds.x, graph_bounds.y, 550, 360), .7, "screen", layout),
    ])
    report: dict = {"passed": False}

    def finished(result: Result | None) -> None:
        try:
            if result is None or result.error:
                raise RuntimeError(result.error if result else "No result")
            assert result.pptx is None
            assert set(output.glob("*.pptx")) == previous_decks
            assert Image.open(result.session / "001.png").getpixel((5, 5)) == (244, 246, 250), "First tab graph did not update"
            assert Image.open(result.session / "002.png").getpixel((5, 5)) == (214, 233, 218), "Second tab graph did not update"
            save_capture(result, output)
            deck = Presentation(result.pptx)
            assert len(deck.slides) == 3
            for slide, size in zip(list(deck.slides)[1:], ((780, 500), (550, 360))):
                assert slide.shapes[1].image.size == size
            report.update(passed=True, pptx=str(result.pptx), session=str(result.session),
                          frozen=bool(getattr(__import__("sys"), "frozen", False)))
        except Exception as exc:
            report["error"] = str(exc)
        finally:
            (output / "self-test.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
            control.destroy()

    def start() -> None:
        desktop.activate()
        ImageGrab.grab(bbox=bounds.box, all_screens=True).save(output / "mock-window.png")
        CaptureProgress(control, settings, desktop, finished)

    control.after(600, start)
    control.after(20000, lambda: (output.joinpath("timeout.txt").write_text("Timed out"), control.destroy()))
    control.mainloop()
    if not report["passed"]:
        raise RuntimeError(str(report))
