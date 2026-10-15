"""Lossless image slides and collision-safe atomic PowerPoint output."""
from __future__ import annotations

import os
import tempfile
from datetime import datetime
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt

from .models import TocGroup


def _page_number(slide: object, number: int, slide_width: int) -> None:
    text = slide.shapes.add_textbox(slide_width - Inches(.95), Inches(7.31), Inches(.75), Inches(.16))
    frame = text.text_frame
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    paragraph = frame.paragraphs[0]
    paragraph.text = str(number)
    paragraph.alignment = PP_ALIGN.RIGHT
    paragraph.font.name, paragraph.font.size = "Malgun Gothic", Pt(10)
    paragraph.font.color.rgb = RGBColor.from_string("555555")


def _toc_rows(captures: list[tuple[str, Path]], groups: list[TocGroup],
              first_content_page: int) -> list[tuple[str, str, bool]]:
    membership = {name: group.name for group in groups for name in group.tabs}
    rows: list[tuple[str, str, bool]] = []
    shown: set[str] = set()
    for index, (name, _) in enumerate(captures):
        group_name = membership.get(name)
        if not group_name:
            rows.append((name, str(first_content_page + index), False))
            continue
        if group_name in shown:
            continue
        shown.add(group_name)
        members = [(member_index, member_name) for member_index, (member_name, _) in enumerate(captures)
                   if membership.get(member_name) == group_name]
        pages = [first_content_page + member_index for member_index, _ in members]
        page_label = f"{pages[0]}–{pages[-1]}" if pages == list(range(pages[0], pages[-1] + 1)) else ", ".join(map(str, pages))
        rows.append((group_name, page_label, True))
        rows.extend((member_name, str(first_content_page + member_index), False) for member_index, member_name in members)
    return rows


def _add_toc_page(prs: Presentation, rows: list[tuple[str, str, bool]], number: int,
                  continued: bool) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(.55), Inches(.48), Inches(11.8), Inches(.65))
    title.text_frame.paragraphs[0].text = "목차" + (" (계속)" if continued else "")
    title.text_frame.paragraphs[0].font.name = "Malgun Gothic"
    title.text_frame.paragraphs[0].font.size = Pt(34)
    title.text_frame.paragraphs[0].font.bold = True
    title.text_frame.paragraphs[0].font.color.rgb = RGBColor.from_string("17324D")
    columns = 1 if len(rows) <= 15 else 2 if len(rows) <= 30 else 3
    column_step = 6.0 if columns <= 2 else 4.2
    label_width = 4.75 if columns <= 2 else 3.05
    for offset, (name, page, heading) in enumerate(rows):
        column, row = divmod(offset, 15)
        x = .65 + column * column_step
        y = 1.43 + row * .37
        label = slide.shapes.add_textbox(Inches(x + (0 if heading else .18)), Inches(y),
                                         Inches(label_width if heading else label_width - .18), Inches(.27))
        frame = label.text_frame
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        frame.word_wrap = False
        p = frame.paragraphs[0]
        p.text = name
        p.font.name, p.font.size = "Malgun Gothic", Pt(15 if heading else 12)
        p.font.bold = heading
        p.font.color.rgb = RGBColor.from_string("17324D" if heading else "333333")
        page_text = slide.shapes.add_textbox(Inches(x + label_width + .1), Inches(y), Inches(.75), Inches(.27))
        page_frame = page_text.text_frame
        page_frame.margin_left = page_frame.margin_right = page_frame.margin_top = page_frame.margin_bottom = 0
        page_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        page_frame.word_wrap = False
        page_paragraph = page_frame.paragraphs[0]
        page_paragraph.text = page
        page_paragraph.alignment = PP_ALIGN.RIGHT
        page_paragraph.font.name, page_paragraph.font.size = "Malgun Gothic", Pt(12)
        page_paragraph.font.color.rgb = RGBColor.from_string("555555")
    _page_number(slide, number, prs.slide_width)


def export_pptx(captures: list[tuple[str, Path]], output: Path, partial: bool = False,
                groups: list[TocGroup] | None = None) -> Path:
    if not captures:
        raise ValueError("저장할 캡처가 없습니다.")
    output.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
    entry_count = len(captures) + sum(1 for group in groups or [] if any(name in group.tabs for name, _ in captures))
    toc_pages = max(1, (entry_count + 44) // 45)
    rows = _toc_rows(captures, groups or [], toc_pages + 1)
    for page in range(toc_pages):
        _add_toc_page(prs, rows[page * 45:(page + 1) * 45], page + 1, page > 0)
    for number, (name, image_path) in enumerate(captures, toc_pages + 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        text = slide.shapes.add_textbox(Inches(.20), Inches(.12), Inches(12.93), Inches(.25))
        text.text_frame.margin_left = text.text_frame.margin_right = 0
        text.text_frame.margin_top = text.text_frame.margin_bottom = 0
        text.text_frame.auto_size = MSO_AUTO_SIZE.NONE
        text.text_frame.word_wrap = False
        p = text.text_frame.paragraphs[0]
        p.text = name
        p.alignment = PP_ALIGN.LEFT
        p.font.name, p.font.size = "Malgun Gothic", Pt(10)
        p.font.color.rgb = RGBColor.from_string("222222")
        picture = slide.shapes.add_picture(str(image_path), 0, 0)
        image_top = text.top + text.height + Inches(.08)
        available_w = Inches(12.93)
        available_h = prs.slide_height - image_top - Inches(.20)
        ratio = min(available_w / picture.width, available_h / picture.height)
        picture.width, picture.height = round(picture.width * ratio), round(picture.height * ratio)
        picture.left = (prs.slide_width - picture.width) // 2
        picture.top = image_top + (available_h - picture.height) // 2
        _page_number(slide, number, prs.slide_width)
    stem = "Bloomberg_" + datetime.now().strftime("%Y%m%d_%H%M%S") + ("_partial" if partial else "")
    fd, raw_path = tempfile.mkstemp(prefix=".bloomberg-", suffix=".tmp", dir=output)
    os.close(fd)
    temporary = Path(raw_path)
    try:
        prs.save(str(temporary))
        for number in range(10000):
            destination = output / f"{stem}{'_' + str(number) if number else ''}.pptx"
            try:
                # Windows rename fails if destination exists: never overwrite another deck.
                temporary.rename(destination)
                return destination
            except FileExistsError:
                continue
        raise OSError("동일한 이름의 파일이 너무 많습니다.")
    finally:
        temporary.unlink(missing_ok=True)
    raise AssertionError("unreachable")
